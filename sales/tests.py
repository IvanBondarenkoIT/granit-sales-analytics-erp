from datetime import date
from decimal import Decimal
from io import BytesIO

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse

from core.models import Product, ProductGroup, Store
from sales.excel_matrix import matrix_to_xlsx
from sales.management.commands.seed_supergroups import ensure_catalog
from sales.matrix import build_matrix
from sales.models import (
    SCOPE_ALL,
    SOURCE_INVOICE,
    SOURCE_REALIZED,
    SOURCE_SERVICE,
    SaleFact,
    SuperGroup,
    SuperGroupMember,
)
from sales.supergroups import load_seed_config, resolve_supergroup_key


class SupergroupResolveTests(TestCase):
    def test_by_group_id_and_prefix(self):
        cfg = load_seed_config()
        self.assertEqual(resolve_supergroup_key(11077, "", config=cfg), "coffee_kg")
        self.assertEqual(resolve_supergroup_key(999, "DRINKS > Water > Still", config=cfg), "water")
        self.assertEqual(resolve_supergroup_key(999, "Cards Loyalty", config=cfg), "exclude")
        self.assertEqual(resolve_supergroup_key(999, "Something else", config=cfg), "unclassified")


class MatrixBalanceTests(TestCase):
    def setUp(self):
        self.store_a = Store.objects.create(granit_id=1, name="Shop A")
        self.store_b = Store.objects.create(granit_id=2, name="Shop B")
        self.g_coffee = ProductGroup.objects.create(granit_id=11077, name="Coffee kg")
        self.g_other = ProductGroup.objects.create(granit_id=99901, name="Misc")
        self.p1 = Product.objects.create(granit_id=101, name="Bean", group=self.g_coffee)
        self.p2 = Product.objects.create(granit_id=102, name="Other", group=self.g_other)
        self.p3 = Product.objects.create(granit_id=103, name="Orphan")  # no group
        coffee = SuperGroup.objects.create(key="coffee_kg", title="coffee kg", color="F4B183", sort_order=1)
        SuperGroup.objects.create(
            key="outside", title="Outside", color="FFFF00", sort_order=10000, is_catchall=True
        )
        SuperGroupMember.objects.create(super_group=coffee, product_group=self.g_coffee)
        SaleFact.objects.create(
            sale_date=date(2026, 9, 21),
            store=self.store_a,
            product=self.p1,
            quantity=Decimal("2"),
            amount=Decimal("100.00"),
            granit_sale_id=1,
            granit_line_id=1,
        )
        SaleFact.objects.create(
            sale_date=date(2026, 9, 22),
            store=self.store_b,
            product=self.p2,
            quantity=Decimal("1"),
            amount=Decimal("30.00"),
            granit_sale_id=2,
            granit_line_id=1,
        )
        SaleFact.objects.create(
            sale_date=date(2026, 9, 22),
            store=self.store_a,
            product=self.p3,
            quantity=Decimal("3"),
            amount=Decimal("15.00"),
            granit_sale_id=3,
            granit_line_id=1,
        )

    def test_matrix_totals_match_facts(self):
        matrix = build_matrix(date(2026, 9, 21), date(2026, 9, 22))
        self.assertTrue(matrix.balanced)
        self.assertEqual(matrix.grand.amount, Decimal("145.00"))
        self.assertEqual(matrix.fact_total.amount, Decimal("145.00"))
        titles = [r.title for r in matrix.rows]
        self.assertIn("coffee kg", titles)
        self.assertTrue(any(r.is_catchall for r in matrix.rows))
        catch = next(r for r in matrix.rows if r.is_catchall)
        self.assertEqual(catch.total.amount, Decimal("45.00"))

    def test_matrix_page_qty_without_trailing_zeros(self):
        SaleFact.objects.create(
            sale_date=date(2026, 9, 21),
            store=self.store_a,
            product=self.p1,
            quantity=Decimal("0.250"),
            amount=Decimal("30.00"),
            granit_sale_id=4,
            granit_line_id=1,
        )
        user = get_user_model().objects.create_user("qty-viewer", password="test-pass-123")
        self.client.force_login(user)
        response = self.client.get(
            reverse("sales_matrix"),
            {"date_from": "2026-09-21", "date_to": "2026-09-22"},
            HTTP_ACCEPT_LANGUAGE="ru",
        )
        self.assertEqual(response.status_code, 200)
        html = response.content.decode()
        self.assertIn("2,25", html)
        self.assertNotIn(",000", html)

    def test_excel_export_bytes(self):
        from openpyxl import load_workbook
        from io import BytesIO

        matrix = build_matrix(date(2026, 9, 21), date(2026, 9, 22))
        data = matrix_to_xlsx(matrix)
        self.assertGreater(len(data), 100)
        self.assertEqual(data[:2], b"PK")
        wb = load_workbook(BytesIO(data))
        self.assertEqual(
            wb.sheetnames,
            ["Продажи", "Supergroups", "Unmapped", "Справочник групп", "Facts"],
        )
        sg = wb["Supergroups"]
        self.assertEqual(sg["A1"].value, "supergroup_key")
        self.assertIn("coffee kg", [c.value for c in sg["B"] if c.value])
        qty_cell = sg.cell(2, 3)
        self.assertTrue(sg.cell(1, 3).value.endswith(" qty"))
        self.assertEqual(qty_cell.number_format, "#,##0.00#")
        amt_header = next(c for c in sg[1] if c.value and str(c.value).endswith(" amt"))
        amt_cell = sg.cell(2, amt_header.column)
        self.assertEqual(amt_cell.number_format, "#,##0.00")
        sales = wb["Продажи"]
        self.assertEqual(sales["A1"].value, "group_id")
        # last row is ИТОГО
        last = sales.max_row
        self.assertEqual(sales.cell(last, 2).value, "ИТОГО")


class SeedAndViewTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user("analyst", password="test-pass-123")
        self.client.force_login(self.user)
        ProductGroup.objects.create(granit_id=11077, name="Coffee kg", parent_path="Coffee kg")
        drinks = ProductGroup.objects.create(granit_id=50000, name="DRINKS")
        ProductGroup.objects.create(granit_id=50001, name="Water", parent=drinks)

    def test_seed_assigns_by_id_and_prefix(self):
        stats = ensure_catalog(fill_unmapped=False)
        self.assertGreater(stats["created_sg"], 0)
        coffee = SuperGroup.objects.get(key="coffee_kg")
        water = SuperGroup.objects.get(key="water")
        self.assertTrue(SuperGroupMember.objects.filter(super_group=coffee, product_group__granit_id=11077).exists())
        self.assertTrue(SuperGroupMember.objects.filter(super_group=water, product_group__granit_id=50001).exists())
        # second plain seed refused by command, but ensure_catalog skips
        again = ensure_catalog(fill_unmapped=False)
        self.assertEqual(again["skipped"], 1)

    def test_matrix_page_ok(self):
        ensure_catalog()
        store = Store.objects.create(granit_id=1, name="Shop")
        product = Product.objects.create(granit_id=1, name="X", group=ProductGroup.objects.get(granit_id=11077))
        SaleFact.objects.create(
            sale_date=date(2026, 9, 21),
            store=store,
            product=product,
            quantity=1,
            amount=10,
            granit_sale_id=9,
            granit_line_id=1,
        )
        response = self.client.get(reverse("sales_matrix"), {"date_from": "2026-09-21", "date_to": "2026-09-21"})
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "coffee kg")

    def test_add_remove_member(self):
        ensure_catalog()
        sg = SuperGroup.objects.get(key="coffee_kg")
        other = ProductGroup.objects.create(granit_id=777, name="Loose")
        response = self.client.post(
            reverse("sales_matrix_sg", args=[sg.pk]) + "?date_from=2026-09-01&date_to=2026-09-30",
            {"action": "add", "group_id": "777"},
        )
        self.assertEqual(response.status_code, 302)
        self.assertTrue(SuperGroupMember.objects.filter(super_group=sg, product_group=other).exists())
        response = self.client.post(
            reverse("sales_matrix_sg", args=[sg.pk]) + "?date_from=2026-09-01&date_to=2026-09-30",
            {"action": "remove", "group_id": "777"},
        )
        self.assertEqual(response.status_code, 302)
        self.assertFalse(SuperGroupMember.objects.filter(product_group=other).exists())

    def test_set_group_owner(self):
        ensure_catalog()
        coffee = SuperGroup.objects.get(key="coffee_kg")
        water = SuperGroup.objects.get(key="water")
        group = ProductGroup.objects.create(granit_id=888, name="Reassign me")
        SuperGroupMember.objects.create(super_group=coffee, product_group=group)
        url = reverse("sales_matrix_group", args=[888]) + "?date_from=2026-09-01&date_to=2026-09-30"
        response = self.client.post(url, {"action": "set_owner", "sg_id": str(water.pk)})
        self.assertEqual(response.status_code, 302)
        self.assertTrue(SuperGroupMember.objects.filter(super_group=water, product_group=group).exists())
        response = self.client.post(url, {"action": "set_owner", "sg_id": "outside"})
        self.assertEqual(response.status_code, 302)
        self.assertFalse(SuperGroupMember.objects.filter(product_group=group).exists())
        page = self.client.get(url)
        self.assertEqual(page.status_code, 200)
        self.assertContains(page, "name=\"sg_id\"")

    def test_move_super_group_order(self):
        ensure_catalog()
        url = reverse("sales_matrix_edit") + "?date_from=2026-09-01&date_to=2026-09-30"
        ordered = list(SuperGroup.objects.filter(is_catchall=False).order_by("sort_order", "title", "pk"))
        self.assertGreaterEqual(len(ordered), 2)
        first, second = ordered[0], ordered[1]
        response = self.client.post(url, {"action": "move_down", "sg_id": str(first.pk)})
        self.assertEqual(response.status_code, 302)
        reordered = list(SuperGroup.objects.filter(is_catchall=False).order_by("sort_order", "title", "pk"))
        self.assertEqual(reordered[0].pk, second.pk)
        self.assertEqual(reordered[1].pk, first.pk)
        response = self.client.post(url, {"action": "move_up", "sg_id": str(first.pk)})
        self.assertEqual(response.status_code, 302)
        restored = list(SuperGroup.objects.filter(is_catchall=False).order_by("sort_order", "title", "pk"))
        self.assertEqual(restored[0].pk, first.pk)


class ScopeTests(TestCase):
    """Retail = receipt by receipt date; all = realized + invoice + service by realization date."""

    def setUp(self):
        self.user = get_user_model().objects.create_user("scope-viewer", password="test-pass-123")
        self.client.force_login(self.user)
        self.shop = Store.objects.create(granit_id=33, name="BatumiMall")
        self.horeca = Store.objects.create(granit_id=178, name="HORECA TBILISI", kind=Store.KIND_WAREHOUSE)
        group = ProductGroup.objects.create(granit_id=11077, name="Coffee kg")
        self.bean = Product.objects.create(granit_id=20575, name="Bean 250 g", group=group)
        common = {"sale_date": date(2026, 8, 14), "product": self.bean, "granit_line_id": 1}
        SaleFact.objects.create(store=self.shop, quantity=Decimal("0.250"), amount=Decimal("30.00"),
                                granit_sale_id=5, **common)
        SaleFact.objects.create(store=self.shop, source=SOURCE_REALIZED, quantity=Decimal("0.250"),
                                amount=Decimal("30.00"), granit_sale_id=5, **common)
        SaleFact.objects.create(store=self.horeca, source=SOURCE_INVOICE, quantity=Decimal("5.000"),
                                amount=Decimal("600.00"), granit_sale_id=5, **common)
        self.period = {"date_from": "2026-08-14", "date_to": "2026-08-14"}

    def test_matrix_retail_unchanged_all_adds_warehouse(self):
        retail = build_matrix(date(2026, 8, 14), date(2026, 8, 14))
        self.assertEqual(retail.grand.amount, Decimal("30.00"))
        self.assertEqual([s.name for s in retail.stores], ["BatumiMall"])
        full = build_matrix(date(2026, 8, 14), date(2026, 8, 14), SCOPE_ALL)
        self.assertTrue(full.balanced)
        self.assertEqual(full.grand.amount, Decimal("630.00"))
        self.assertEqual([s.name for s in full.stores], ["BatumiMall", "HORECA TBILISI"])
        self.assertIn("scope=all", full.period_query)

    def test_pages_follow_scope(self):
        page = self.client.get(reverse("sales_matrix"), self.period)
        self.assertNotContains(page, "HORECA TBILISI</th>")
        page = self.client.get(reverse("sales_matrix"), {**self.period, "scope": "all"})
        self.assertContains(page, "HORECA TBILISI</th>")
        self.assertContains(page, "scope=all")

        explorer = self.client.get(reverse("sales_explorer"), {**self.period, "group_by": "receipt"})
        self.assertEqual(explorer.context["total_amount"], Decimal("30.00"))
        explorer = self.client.get(
            reverse("sales_explorer"), {**self.period, "group_by": "receipt", "scope": "all"}
        )
        self.assertEqual(explorer.context["total_amount"], Decimal("630.00"))
        self.assertContains(explorer, "source=invoice")

        group = self.client.get(reverse("sales_matrix_group", args=[11077]), {**self.period, "scope": "all"})
        self.assertEqual(group.context["total_amount"], Decimal("630.00"))

    def test_service_group_with_negative_id(self):
        services = ProductGroup.objects.create(granit_id=-1, name="Услуги")
        repair = Product.objects.create(granit_id=-1, name="Услуга: ремонт", group=services)
        SaleFact.objects.create(sale_date=date(2026, 8, 14), store=self.horeca, product=repair,
                                source=SOURCE_SERVICE, quantity=Decimal("1"), amount=Decimal("50.00"),
                                granit_sale_id=6, granit_line_id=1)
        outside = SuperGroup.objects.create(key="outside", title="Outside", is_catchall=True)
        page = self.client.get(reverse("sales_matrix_sg", args=[outside.pk]), {**self.period, "scope": "all"})
        self.assertContains(page, "/sales/matrix/group/-1/")
        group = self.client.get(reverse("sales_matrix_group", args=[-1]), {**self.period, "scope": "all"})
        self.assertEqual(group.status_code, 200)
        self.assertEqual(group.context["total_amount"], Decimal("50.00"))

    def test_receipt_and_invoice_with_same_id(self):
        receipt = self.client.get(reverse("sales_receipt", args=[5]))
        self.assertEqual(receipt.context["total_amount"], Decimal("30.00"))
        invoice = self.client.get(reverse("sales_receipt", args=[5]), {"source": "invoice"})
        self.assertEqual(invoice.context["total_amount"], Decimal("600.00"))
        self.assertTrue(invoice.context["is_invoice"])

    def test_realized_drill_down_from_all_scope(self):
        SaleFact.objects.filter(source=SOURCE_REALIZED, granit_sale_id=5).update(amount=Decimal("29.670330"))
        explorer = self.client.get(
            reverse("sales_explorer"), {**self.period, "group_by": "receipt", "scope": "all"}
        )
        self.assertContains(explorer, "source=realized")
        page = self.client.get(reverse("sales_receipt", args=[5]), {"source": "realized"})
        self.assertEqual(page.context["total_amount"], Decimal("29.670330"))
        self.assertTrue(page.context["is_realized"])
        self.assertContains(page, "29,67")

    def test_promo_analysis_stays_retail(self):
        from promos.analysis import _daily_totals, _warehouse_span

        totals = _daily_totals([self.bean.id], date(2026, 8, 14), date(2026, 8, 14))
        self.assertEqual(totals, {date(2026, 8, 14): Decimal("30.00")})
        # Only retail bounds: invoice/realized on another day must not widen the span.
        SaleFact.objects.create(
            store=self.horeca,
            product=self.bean,
            source=SOURCE_INVOICE,
            sale_date=date(2026, 1, 1),
            quantity=Decimal("1"),
            amount=Decimal("10"),
            granit_sale_id=99,
            granit_line_id=1,
        )
        first, last = _warehouse_span()
        self.assertEqual(first, date(2026, 8, 14))
        self.assertEqual(last, date(2026, 8, 14))


class ReconcilePageTests(TestCase):
    def setUp(self):
        from etl.models import ReconcileSnapshot

        self.user = get_user_model().objects.create_user("viewer", password="test-pass-123")
        self.client.force_login(self.user)
        self.store = Store.objects.create(granit_id=133, name="СЦ ГРУЗИЯ")
        store_row = {
            "store": 133, "name": "СЦ ГРУЗИЯ", "docs": 3,
            "warehouses": [{"id": 133, "name": "СЦ ГРУЗИЯ", "amount": 150.0}],
            "granit_total": 150.0, "granit_goods": 40.0, "granit_services": 100.0, "leftover": 10.0,
            "ours_goods": 40.0, "ours_services": 100.0, "ours_total": 140.0,
            "diff": -10.0, "diff_goods": 0.0, "diff_services": 0.0, "flags": ["leftover", "diff"],
        }
        ReconcileSnapshot.objects.create(
            month=date(2026, 8, 1),
            granit_total=Decimal("150"),
            ours_total=Decimal("140"),
            flag_count=2,
            data={
                "month": "2026-08-01",
                "stores": [store_row],
                "totals": {k: v for k, v in store_row.items() if isinstance(v, float)},
                "services": [{"id": 1, "name": "ремонт кофеварки", "granit": 100.0, "ours": 100.0, "new": True}],
                "other_types": [{"typ": 1, "docs": 4, "amount": 500.0}],
                "flags": [
                    {"kind": "leftover", "store": "СЦ ГРУЗИЯ", "amount": 10.0},
                    {"kind": "diff", "store": "СЦ ГРУЗИЯ", "amount": -10.0},
                ],
            },
        )

    def test_overview_month_and_matrix_badge(self):
        page = self.client.get(reverse("sales_reconcile"))
        self.assertEqual(page.status_code, 200)
        self.assertEqual(page.context["current"]["status"], "bad")
        self.assertEqual(page.context["services"]["kinds"][0]["new_in"], "2026-08")
        self.assertContains(page, reverse("sales_reconcile_month", args=["2026-08"]))

        month = self.client.get(reverse("sales_reconcile_month", args=["2026-08"]))
        self.assertEqual(month.context["stores"][0]["status"], "bad")
        self.assertContains(month, reverse("sales_reconcile_docs", args=["2026-08", 133]))
        self.assertEqual(self.client.get(reverse("sales_reconcile_month", args=["2026-07"])).status_code, 404)

        matrix = self.client.get(reverse("sales_matrix"), {"date_from": "2026-08-01", "date_to": "2026-08-31"})
        self.assertEqual(matrix.context["reconcile"]["key"], "2026-08")
        self.assertContains(matrix, "rc-badge bad")

    def test_open_month_is_not_a_mismatch(self):
        from etl.models import ReconcileSnapshot

        snap = ReconcileSnapshot.objects.get(month=date(2026, 8, 1))
        snap.data["state"] = "open"
        snap.save()
        page = self.client.get(reverse("sales_reconcile"), HTTP_ACCEPT_LANGUAGE="ru")
        self.assertEqual(page.context["current"]["status"], "open")
        month = self.client.get(reverse("sales_reconcile_month", args=["2026-08"]), HTTP_ACCEPT_LANGUAGE="ru")
        self.assertContains(month, "Месяц не закрыт")
        matrix = self.client.get(reverse("sales_matrix"), {"date_from": "2026-08-01", "date_to": "2026-08-31"})
        self.assertContains(matrix, "rc-badge open")

    def test_documents_page(self):
        from unittest.mock import patch

        rows = [{
            "kind": "invoice", "key": 10, "docs": [{"id": 10, "nbr": "7", "doc_date": date(2026, 8, 3),
                                                   "real_date": date(2026, 8, 3)}],
            "granit": 150.0, "goods": 40.0, "services": 100.0, "leftover": 10.0,
            "ours": 140.0, "ours_dates": [date(2026, 8, 3)], "diff": -10.0,
        }]
        with patch("sales.reconcile_views.month_documents", return_value={"rows": rows, "total": 1, "checked": 5}):
            page = self.client.get(reverse("sales_reconcile_docs", args=["2026-08", 133]))
        self.assertEqual(page.status_code, 200)
        self.assertContains(page, "DGVDT 10")

    def test_refresh_is_staff_only(self):
        from unittest.mock import patch

        url = reverse("sales_reconcile_refresh")
        with patch("sales.reconcile_views.refresh_months") as refresh:
            self.assertEqual(self.client.post(url).status_code, 403)
            refresh.assert_not_called()
            self.user.is_staff = True
            self.user.save()
            response = self.client.post(url, {"next": "https://evil.example/"})
            refresh.assert_called_once()
        self.assertEqual(response["Location"], reverse("sales_reconcile"))


class BlaserCertTests(TestCase):
    def setUp(self):
        from core.models import Client
        from sales.blaser_cert import BLASER_GROUP_IDS

        self.store = Store.objects.create(granit_id=10, name="Mall")
        self.g250 = ProductGroup.objects.create(granit_id=BLASER_GROUP_IDS[0], name="Blaser 250g")
        self.g1kg = ProductGroup.objects.create(granit_id=BLASER_GROUP_IDS[1], name="Blaser 1kg")
        self.g_maker = ProductGroup.objects.create(granit_id=21270, name="Delonghi coffee makers")
        self.bean = Product.objects.create(granit_id=501, name="Blend 250", group=self.g250)
        self.maker = Product.objects.create(granit_id=601, name="Magnifica S", group=self.g_maker)
        makers = SuperGroup.objects.create(key="coffee_makers", title="coffee makers", color="AABBCC")
        SuperGroupMember.objects.create(super_group=makers, product_group=self.g_maker)
        self.alice = Client.objects.create(granit_id=1001, name="Alice", phone="+995555000001", card_number="2200000093516")
        self.bob = Client.objects.create(granit_id=1002, name="Bob", phone="+995555000002")
        self.cara = Client.objects.create(granit_id=1003, name="Cara")

        def sale(day, client, product, qty, amount, sale_id, line_id):
            SaleFact.objects.create(
                sale_date=day,
                store=self.store,
                client=client,
                product=product,
                quantity=Decimal(qty),
                amount=Decimal(amount),
                granit_sale_id=sale_id,
                granit_line_id=line_id,
            )

        # Full price 50 dominates; 45 = 10% off.
        for i, day in enumerate([date(2025, 11, 1), date(2025, 11, 2), date(2025, 11, 3), date(2025, 11, 4)]):
            sale(day, self.cara, self.bean, "1", "50.00", 10 + i, 1)
        sale(date(2025, 12, 1), self.alice, self.bean, "1", "45.00", 20, 1)
        sale(date(2026, 1, 10), self.alice, self.bean, "2", "90.00", 21, 1)  # second visit at 10%
        sale(date(2025, 12, 15), self.bob, self.bean, "1", "45.00", 22, 1)  # once
        sale(date(2025, 12, 2), self.alice, self.bean, "1", "40.00", 23, 1)  # other price
        sale(date(2025, 10, 20), self.alice, self.maker, "1", "1200.00", 30, 1)
        # Invoice machine for Bob (must count in machine_cross / enrichment).
        SaleFact.objects.create(
            sale_date=date(2025, 11, 5),
            source=SOURCE_INVOICE,
            store=self.store,
            client=self.bob,
            product=self.maker,
            quantity=Decimal("1"),
            amount=Decimal("1500.00"),
            granit_sale_id=31,
            granit_line_id=1,
        )

    def test_calibrate_and_classify(self):
        from sales.blaser_cert import calibrate_full_price, classify_price

        full = calibrate_full_price([Decimal("50")] * 5 + [Decimal("45")] * 3)
        self.assertEqual(full, Decimal("50.00"))
        self.assertEqual(classify_price(Decimal("50.00"), full), "full")
        self.assertEqual(classify_price(Decimal("45.00"), full), "promo_10")
        self.assertEqual(classify_price(Decimal("40.00"), full), "other")

    def test_report_once_multi_and_machine(self):
        from sales.blaser_cert import build_report, write_report_xlsx
        from openpyxl import load_workbook
        from pathlib import Path
        import tempfile

        report = build_report(date(2025, 10, 15), date(2026, 10, 16))
        self.assertEqual(report.summary["promo_clients"], 2)
        self.assertEqual(report.summary["visits_once"], 1)
        self.assertEqual(report.summary["visits_multi"], 1)
        by_id = {r.granit_id: r for r in report.clients}
        self.assertEqual(by_id[1001].visits, 2)
        self.assertEqual(by_id[1001].machine_count, 1)
        self.assertIn("Magnifica S", by_id[1001].machines)
        self.assertEqual(by_id[1002].visits, 1)
        self.assertEqual(by_id[1002].machine_count, 1)
        self.assertIn("Magnifica S", by_id[1002].machines)
        self.assertEqual(report.summary["machine_buyers_in_window"], 2)
        self.assertEqual(report.summary["machine_buyers_used_promo"], 2)

        with tempfile.TemporaryDirectory() as tmp:
            path = write_report_xlsx(report, Path(tmp) / "out.xlsx")
            wb = load_workbook(path)
            self.assertIn("summary", wb.sheetnames)
            self.assertIn("clients", wb.sheetnames)
            self.assertIn("sms_ready", wb.sheetnames)
            self.assertGreaterEqual(wb["sms_ready"].max_row, 3)