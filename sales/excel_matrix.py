"""Excel export matching monthly-sales-report sales_matrix layout (5 sheets)."""
from __future__ import annotations

from collections import defaultdict
from datetime import date
from decimal import Decimal
from io import BytesIO

from django.db.models import Count, Sum
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from core.models import Product, ProductGroup, Store
from sales.matrix import MatrixResult, ZERO, ensure_catchall, sales_in_period
from sales.models import SuperGroup, SuperGroupMember

TOTALS_LABEL = "ИТОГО"
QTY_FORMAT = "#,##0.00#"


def _fill(color: str) -> PatternFill:
    hex_color = (color or "FFFFFF").lstrip("#").upper()[:6].rjust(6, "0")
    return PatternFill(start_color=hex_color, end_color=hex_color, fill_type="solid")


def _store_prefix(store: Store) -> str:
    return f"{store.name} ({store.granit_id})"


def _sg_lookup() -> tuple[dict[int, SuperGroup], SuperGroup]:
    """product_group_id -> SuperGroup; catchall for unmapped."""
    catchall = ensure_catchall()
    members = {
        m.product_group_id: m.super_group
        for m in SuperGroupMember.objects.select_related("super_group")
        if not m.super_group.is_catchall
    }
    return members, catchall


def _dominant_param_by_group(group_ids: list[int]) -> dict[int, str]:
    if not group_ids:
        return {}
    rows = (
        Product.objects.filter(group_id__in=group_ids, parameter_value__isnull=False)
        .values("group_id", "parameter_value__label")
        .annotate(n=Count("id"))
        .order_by("group_id", "-n")
    )
    out: dict[int, str] = {}
    for row in rows:
        gid = row["group_id"]
        if gid not in out:
            out[gid] = row["parameter_value__label"] or ""
    return out


def _write_header(ws, headers: list[str]) -> None:
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(1, col, header)
        cell.font = Font(bold=True)
        if header != "supergroup_key":
            ws.column_dimensions[get_column_letter(col)].width = max(len(str(header)) + 2, 10)


def _hide_key_col(ws, headers: list[str]) -> None:
    if "supergroup_key" in headers:
        idx = headers.index("supergroup_key") + 1
        ws.column_dimensions[get_column_letter(idx)].hidden = True


def _style_metric_cell(cell, header: str) -> None:
    if header.endswith(" qty") or header in {"qty", "total_qty"}:
        cell.number_format = QTY_FORMAT
    elif header.endswith(" amt") or header in {"amount", "total_amount"}:
        cell.number_format = "#,##0"
    elif header == "group_id" and cell.value not in ("", None):
        cell.number_format = "0"


def _append_row(ws, values: list, headers: list[str], *, fill: PatternFill | None = None, bold: bool = False) -> None:
    ws.append(values)
    row_idx = ws.max_row
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row_idx, col)
        if bold:
            cell.font = Font(bold=True)
        if fill is not None and not bold:
            cell.fill = fill
        _style_metric_cell(cell, header)


def _freeze_and_filter(ws, headers: list[str], freeze_cols: int, *, exclude_last: bool) -> None:
    if freeze_cols > 0:
        ws.freeze_panes = f"{get_column_letter(freeze_cols + 1)}2"
    last = ws.max_row
    if exclude_last and last > 1:
        last -= 1
    if headers and last >= 1:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{last}"


def _facts_agg(date_from: date, date_to: date, scope: str):
    """Return long facts: store_id(pk), group_id(pk|None), qty, amount."""
    return list(
        sales_in_period(date_from, date_to, scope)
        .values("store_id", "product__group_id")
        .annotate(qty=Sum("quantity"), amount=Sum("amount"))
    )


def matrix_to_xlsx(matrix: MatrixResult) -> bytes:
    date_from, date_to = matrix.date_from, matrix.date_to
    stores = matrix.stores
    members, catchall = _sg_lookup()
    colors = {sg.key: sg.color for sg in SuperGroup.objects.all()}
    colors.setdefault(catchall.key, catchall.color)

    by_store = {s.id: s for s in stores}
    groups = {g.id: g for g in ProductGroup.objects.all()}
    facts = _facts_agg(date_from, date_to, matrix.scope)
    sold_group_ids = {f["product__group_id"] for f in facts if f["product__group_id"] is not None}
    params = _dominant_param_by_group(list(sold_group_ids))

    # --- aggregate group × store ---
    group_store: dict[int | None, dict[int, list[Decimal]]] = defaultdict(lambda: defaultdict(lambda: [ZERO, ZERO]))
    for f in facts:
        gid = f["product__group_id"]
        sid = f["store_id"]
        group_store[gid][sid][0] += f["qty"] or ZERO
        group_store[gid][sid][1] += f["amount"] or ZERO

    def resolve_sg(group_pk: int | None) -> tuple[str, str]:
        if group_pk is None:
            return catchall.key, catchall.title
        sg = members.get(group_pk)
        if sg is None:
            return catchall.key, catchall.title
        return sg.key, sg.title

    wb = Workbook()
    wb.properties.title = f"Sales {date_from:%Y-%m-%d}..{date_to:%Y-%m-%d} ({matrix.scope})"

    # ===== Продажи =====
    ws_sales = wb.active
    ws_sales.title = "Продажи"
    sales_headers = [
        "group_id",
        "group_name",
        "parent_path",
        "param",
        "supergroup",
        "supergroup_key",
    ]
    for store in stores:
        prefix = _store_prefix(store)
        sales_headers.append(f"{prefix} qty")
        sales_headers.append(f"{prefix} amt")
    sales_headers.extend(["total_qty", "total_amount"])
    _write_header(ws_sales, sales_headers)

    sales_rows_sorted = sorted(
        group_store.keys(),
        key=lambda gid: (
            (groups[gid].name if gid in groups else ""),
            groups[gid].granit_id if gid in groups else 0,
        ),
    )
    for gid in sales_rows_sorted:
        g = groups.get(gid) if gid is not None else None
        key, title = resolve_sg(gid)
        store_map = group_store[gid]
        total_q = sum((v[0] for v in store_map.values()), ZERO)
        total_a = sum((v[1] for v in store_map.values()), ZERO)
        values = [
            g.granit_id if g else "",
            g.name if g else "(no group)",
            g.parent_path if g else "",
            params.get(gid, "") if gid is not None else "",
            title,
            key,
        ]
        for store in stores:
            q, a = store_map.get(store.id, [ZERO, ZERO])
            values.append(float(q))
            values.append(float(a))
        values.append(float(total_q))
        values.append(float(total_a))
        _append_row(ws_sales, values, sales_headers, fill=_fill(colors.get(key, "FFFFFF")))

    # totals
    tot = [""] * 6
    tot[1] = TOTALS_LABEL
    for store in stores:
        col_q = ZERO
        col_a = ZERO
        for store_map in group_store.values():
            q, a = store_map.get(store.id, [ZERO, ZERO])
            col_q += q
            col_a += a
        tot.append(float(col_q))
        tot.append(float(col_a))
    tot.append(float(matrix.grand.qty))
    tot.append(float(matrix.grand.amount))
    _append_row(ws_sales, tot, sales_headers, bold=True)
    _hide_key_col(ws_sales, sales_headers)
    _freeze_and_filter(ws_sales, sales_headers, 5, exclude_last=True)

    # ===== Supergroups =====
    ws_sg = wb.create_sheet("Supergroups")
    sg_headers = ["supergroup_key", "supergroup"]
    for store in stores:
        prefix = _store_prefix(store)
        sg_headers.append(f"{prefix} qty")
        sg_headers.append(f"{prefix} amt")
    sg_headers.extend(["total_qty", "total_amount"])
    _write_header(ws_sg, sg_headers)

    for row in matrix.rows:
        values = [row.key, row.title]
        for store in stores:
            cell = row.cell_for(store.id)
            values.append(float(cell.qty))
            values.append(float(cell.amount))
        values.append(float(row.total.qty))
        values.append(float(row.total.amount))
        _append_row(ws_sg, values, sg_headers, fill=_fill(row.color))

    sg_tot = ["", TOTALS_LABEL]
    for store in stores:
        c = matrix.column_totals[store.id]
        sg_tot.append(float(c.qty))
        sg_tot.append(float(c.amount))
    sg_tot.append(float(matrix.grand.qty))
    sg_tot.append(float(matrix.grand.amount))
    _append_row(ws_sg, sg_tot, sg_headers, bold=True)
    _hide_key_col(ws_sg, sg_headers)
    _freeze_and_filter(ws_sg, sg_headers, 2, exclude_last=True)

    # ===== Unmapped =====
    ws_un = wb.create_sheet("Unmapped")
    un_headers = [
        "group_id",
        "group_name",
        "parent_path",
        "param",
        "supergroup",
        "store_id",
        "store_name",
        "qty",
        "amount",
    ]
    _write_header(ws_un, un_headers)
    for f in facts:
        gid = f["product__group_id"]
        key, title = resolve_sg(gid)
        if key != catchall.key:
            continue
        g = groups.get(gid) if gid is not None else None
        store = by_store.get(f["store_id"])
        if store is None:
            continue
        values = [
            g.granit_id if g else "",
            g.name if g else "(no group)",
            g.parent_path if g else "",
            params.get(gid, "") if gid is not None else "",
            title,
            store.granit_id,
            store.name,
            float(f["qty"] or ZERO),
            float(f["amount"] or ZERO),
        ]
        _append_row(ws_un, values, un_headers, fill=_fill(catchall.color))
    _freeze_and_filter(ws_un, un_headers, 2, exclude_last=False)

    # ===== Справочник групп =====
    ws_cat = wb.create_sheet("Справочник групп")
    cat_headers = [
        "group_id",
        "group_name",
        "parent_path",
        "supergroup",
        "supergroup_key",
        "had_sales",
    ]
    _write_header(ws_cat, cat_headers)
    for g in ProductGroup.objects.order_by("name"):
        key, title = resolve_sg(g.id)
        values = [
            g.granit_id,
            g.name,
            g.parent_path,
            title,
            key,
            1 if g.id in sold_group_ids else 0,
        ]
        _append_row(ws_cat, values, cat_headers, fill=_fill(colors.get(key, "FFFFFF")))
    _hide_key_col(ws_cat, cat_headers)
    _freeze_and_filter(ws_cat, cat_headers, 1, exclude_last=False)

    # ===== Facts =====
    ws_facts = wb.create_sheet("Facts")
    fact_headers = [
        "store_id",
        "store_name",
        "group_id",
        "group_name",
        "parent_path",
        "param",
        "supergroup",
        "supergroup_key",
        "qty",
        "amount",
    ]
    _write_header(ws_facts, fact_headers)
    for f in sorted(facts, key=lambda r: (r["store_id"] or 0, r["product__group_id"] or 0)):
        gid = f["product__group_id"]
        key, title = resolve_sg(gid)
        g = groups.get(gid) if gid is not None else None
        store = by_store.get(f["store_id"])
        if store is None:
            continue
        values = [
            store.granit_id,
            store.name,
            g.granit_id if g else "",
            g.name if g else "(no group)",
            g.parent_path if g else "",
            params.get(gid, "") if gid is not None else "",
            title,
            key,
            float(f["qty"] or ZERO),
            float(f["amount"] or ZERO),
        ]
        _append_row(ws_facts, values, fact_headers, fill=_fill(colors.get(key, "FFFFFF")))
    _hide_key_col(ws_facts, fact_headers)
    _freeze_and_filter(ws_facts, fact_headers, 2, exclude_last=False)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
