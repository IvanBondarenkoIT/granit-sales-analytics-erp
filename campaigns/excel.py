"""Parse SMS campaign Excel (local only; do not commit files)."""
from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import IO

from openpyxl import load_workbook

PHONE_SHEET_NAMES = ("Телефоны", "Phones", "phones")


@dataclass(frozen=True)
class CampaignRow:
    granit_client_id: int
    phone: str


def normalize_phone(raw: object) -> str:
    digits = "".join(ch for ch in str(raw or "") if ch.isdigit())
    if digits.startswith("995") and len(digits) >= 12:
        digits = digits[3:]
    if digits.startswith("0") and len(digits) == 10:
        digits = digits[1:]
    return digits


def _header_map(row: tuple) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for idx, cell in enumerate(row):
        key = str(cell or "").strip().lower()
        if key:
            mapping[key] = idx
    return mapping


def parse_campaign_xlsx(source: str | Path | IO[bytes]) -> list[CampaignRow]:
    wb = load_workbook(source, read_only=True, data_only=True)
    try:
        sheet = None
        for name in wb.sheetnames:
            if name in PHONE_SHEET_NAMES:
                sheet = wb[name]
                break
        if sheet is None:
            sheet = wb[wb.sheetnames[0]]
        rows_iter = sheet.iter_rows(values_only=True)
        header = next(rows_iter, None)
        if not header:
            return []
        cols = _header_map(header)
        id_col = cols.get("customer_id")
        if id_col is None:
            id_col = cols.get("granit_client_id") or cols.get("client_id") or cols.get("id")
        phone_col = cols.get("phone")
        if id_col is None:
            raise ValueError("Excel must have a customer_id column")
        seen: set[int] = set()
        out: list[CampaignRow] = []
        for row in rows_iter:
            if not row or id_col >= len(row):
                continue
            try:
                client_id = int(row[id_col])
            except (TypeError, ValueError):
                continue
            if client_id in seen:
                continue
            seen.add(client_id)
            phone = ""
            if phone_col is not None and phone_col < len(row):
                phone = normalize_phone(row[phone_col])
            out.append(CampaignRow(granit_client_id=client_id, phone=phone))
        return out
    finally:
        wb.close()
