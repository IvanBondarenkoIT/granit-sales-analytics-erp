"""Parse promo Excel (SKU or group). Example lives in docs/; local copy is gitignored."""
from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import IO

from django.conf import settings
from openpyxl import Workbook, load_workbook

PROMO_SHEET_NAMES = ("Акции", "Promos", "promos", "Акция")
LOCAL_PROMOS_XLSX = Path(settings.BASE_DIR) / "data" / "local" / "promos_example.xlsx"
DOCS_PROMOS_XLSX = Path(settings.BASE_DIR) / "docs" / "promos_example.xlsx"

EXAMPLE_HEADERS = [
    "name",
    "start_date",
    "end_date",
    "pre_period_days",
    "group_id",
    "product_id",
    "notes",
]


@dataclass(frozen=True)
class PromoFileRow:
    name: str
    start_date: date | None
    end_date: date | None
    pre_period_days: int | None
    group_granit_id: int | None
    product_granit_id: int | None
    notes: str


def _header_map(row: tuple) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for idx, cell in enumerate(row):
        key = str(cell or "").strip().lower()
        if key:
            mapping[key] = idx
    return mapping


def _col(cols: dict[str, int], *names: str) -> int | None:
    for name in names:
        if name in cols:
            return cols[name]
    return None


def _as_date(raw: object) -> date | None:
    if raw is None or raw == "":
        return None
    if isinstance(raw, datetime):
        return raw.date()
    if isinstance(raw, date):
        return raw
    text = str(raw).strip()[:10]
    try:
        return date.fromisoformat(text)
    except ValueError:
        return None


def _as_int(raw: object) -> int | None:
    if raw is None or raw == "":
        return None
    try:
        return int(raw)
    except (TypeError, ValueError):
        return None


def parse_promo_xlsx(source: str | Path | IO[bytes]) -> list[PromoFileRow]:
    wb = load_workbook(source, read_only=True, data_only=True)
    try:
        sheet = None
        for name in wb.sheetnames:
            if name in PROMO_SHEET_NAMES:
                sheet = wb[name]
                break
        if sheet is None:
            sheet = wb[wb.sheetnames[0]]
        rows_iter = sheet.iter_rows(values_only=True)
        header = next(rows_iter, None)
        if not header:
            return []
        cols = _header_map(header)
        name_col = _col(cols, "name", "название", "promo", "акция")
        start_col = _col(cols, "start_date", "start", "дата_с", "from")
        end_col = _col(cols, "end_date", "end", "дата_по", "to")
        days_col = _col(cols, "pre_period_days", "дни_базы", "baseline_days")
        group_col = _col(cols, "group_id", "группа_id", "group")
        product_col = _col(cols, "product_id", "товар_id", "sku", "product")
        notes_col = _col(cols, "notes", "заметки")
        if group_col is None and product_col is None:
            raise ValueError("Excel must have group_id or product_id")
        out: list[PromoFileRow] = []
        for row in rows_iter:
            if not row:
                continue
            group_id = _as_int(row[group_col]) if group_col is not None and group_col < len(row) else None
            product_id = _as_int(row[product_col]) if product_col is not None and product_col < len(row) else None
            if group_id is None and product_id is None:
                continue
            name = ""
            if name_col is not None and name_col < len(row) and row[name_col]:
                name = str(row[name_col]).strip()
            notes = ""
            if notes_col is not None and notes_col < len(row) and row[notes_col]:
                notes = str(row[notes_col]).strip()
            out.append(
                PromoFileRow(
                    name=name,
                    start_date=_as_date(row[start_col]) if start_col is not None and start_col < len(row) else None,
                    end_date=_as_date(row[end_col]) if end_col is not None and end_col < len(row) else None,
                    pre_period_days=_as_int(row[days_col]) if days_col is not None and days_col < len(row) else None,
                    group_granit_id=group_id,
                    product_granit_id=product_id,
                    notes=notes,
                )
            )
        return out
    finally:
        wb.close()


def write_example_xlsx(path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Акции"
    ws.append(EXAMPLE_HEADERS)
    ws.append(
        [
            "Candies Sep 2026",
            date(2026, 9, 8),
            date(2026, 9, 21),
            14,
            26232,
            None,
            "Example: product group Candies",
        ]
    )
    ws.append(
        [
            "Ballerina Sep 2026",
            date(2026, 9, 8),
            date(2026, 9, 21),
            14,
            None,
            20636,
            'Example: SKU Coffee "Blaser" Ballerina (250 g)',
        ]
    )
    wb.save(path)
    return path


def ensure_example_files() -> Path | None:
    if not DOCS_PROMOS_XLSX.is_file():
        write_example_xlsx(DOCS_PROMOS_XLSX)
    if not LOCAL_PROMOS_XLSX.is_file():
        write_example_xlsx(LOCAL_PROMOS_XLSX)
    if LOCAL_PROMOS_XLSX.is_file():
        return LOCAL_PROMOS_XLSX
    if DOCS_PROMOS_XLSX.is_file():
        return DOCS_PROMOS_XLSX
    return None
